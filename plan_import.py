"""解析輔導上傳的 xlsx 讀經計畫：date / book / range / guiding_question 四欄
（跟天父日記的 plan.xlsx 同一種欄位，方便你把既有的計畫表拿來改一改就能用）。

`guiding_question` 欄位可以留空，留空就等那一天真的第一次被打開時才用 AI 生
一句（見 data_store._resolve_guiding_question）——匯入的當下不會呼叫 AI。
"""

from datetime import date, datetime

import openpyxl

from scripture import get_scripture

REQUIRED_COLUMNS = {"date", "book", "range"}


def parse_plan_file(file_stream) -> tuple[list[dict], list[str]]:
    """回傳 (可以排的列, 錯誤訊息)。單一列格式不對只跳過那一列，不會擋掉整批。
    整個檔案包在最外層 try/except——上傳的檔案什麼奇怪格式都可能出現，
    這裡的原則是「解析失敗就回報錯誤訊息」，絕不能讓一個爛檔案把整個網站弄成 500。
    """
    try:
        return _parse_plan_file(file_stream)
    except Exception as exc:  # noqa: BLE001 - 見上方說明
        return [], [f"讀取失敗：{exc}"]


def _parse_plan_file(file_stream) -> tuple[list[dict], list[str]]:
    workbook = openpyxl.load_workbook(file_stream, data_only=True)

    sheet = workbook.active
    if sheet is None:
        return [], ["這個檔案裡沒有工作表"]

    header_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
    if header_row is None:
        return [], ["空白檔案，沒有任何欄位"]

    header = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in header_row]
    missing = REQUIRED_COLUMNS - set(header)
    if missing:
        return [], [f"缺少欄位：{'、'.join(sorted(missing))}（需要 date / book / range，guiding_question 選填）"]

    col_index = {name: i for i, name in enumerate(header)}
    max_index = max(col_index.values())
    rows: list[dict] = []
    errors: list[str] = []

    for row_num, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = [cell.value for cell in row]
        if len(values) <= max_index:
            values.extend([None] * (max_index + 1 - len(values)))
        if all(v is None or str(v).strip() == "" for v in values):
            continue

        try:
            raw_date = values[col_index["date"]]
            book = str(values[col_index["book"]] or "").strip()
            rng = str(values[col_index["range"]] or "").strip()
            guiding_question = ""
            if "guiding_question" in col_index:
                guiding_question = str(values[col_index["guiding_question"]] or "").strip()

            if not raw_date or not book or not rng:
                errors.append(f"第 {row_num} 列缺少 date / book / range，跳過")
                continue

            if isinstance(raw_date, (datetime, date)):
                date_str = raw_date.strftime("%Y-%m-%d")
            else:
                date_str = str(raw_date).strip()
                try:
                    date_str = datetime.strptime(date_str, "%Y-%m-%d").strftime("%Y-%m-%d")
                except ValueError:
                    errors.append(f"第 {row_num} 列的日期「{date_str}」看不懂，格式要 YYYY-MM-DD，跳過")
                    continue

            verses = get_scripture(book, rng)
            if not verses:
                errors.append(f"第 {row_num} 列（{date_str} {book} {rng}）找不到經文，跳過")
                continue

            rows.append(
                {
                    "date": date_str,
                    "reference": f"{book} {rng}",
                    "verses": verses,
                    "guiding_question": guiding_question,
                }
            )
        except Exception as exc:  # noqa: BLE001 - 單一列壞掉不能拖垮整批
            errors.append(f"第 {row_num} 列讀取失敗，跳過：{exc}")
            continue

    return rows, errors
